"""Create the SupportDesk capstone slide deck.

Run from the repository root after installing demo/requirements.txt:

    python demo/create_slides.py
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "supportdesk-slides.pptx"

NAVY = RGBColor(16, 24, 40)
INK = RGBColor(24, 32, 51)
SLATE = RGBColor(83, 96, 120)
MUTED = RGBColor(138, 149, 167)
LINE = RGBColor(226, 231, 239)
PANEL = RGBColor(247, 249, 252)
WHITE = RGBColor(255, 255, 255)
BLUE = RGBColor(57, 112, 188)
BLUE_LIGHT = RGBColor(232, 243, 255)
GREEN = RGBColor(45, 137, 101)
GREEN_LIGHT = RGBColor(232, 248, 240)
AMBER = RGBColor(186, 122, 29)
AMBER_LIGHT = RGBColor(255, 246, 224)
PURPLE = RGBColor(119, 82, 160)
PURPLE_LIGHT = RGBColor(244, 237, 255)


def set_background(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False,
             align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_rich_text(slide, lines, x, y, w, h, size=16, color=SLATE,
                  line_gap=4):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    for index, item in enumerate(lines):
        if isinstance(item, str):
            label, text, item_color, item_bold = "", item, color, False
        else:
            label, text, item_color, item_bold = item
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.space_after = Pt(line_gap)
        paragraph.level = 0
        if label:
            run = paragraph.add_run()
            run.text = label
            run.font.name = "Aptos"
            run.font.size = Pt(size)
            run.font.bold = True
            run.font.color.rgb = item_color
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.color.rgb = item_color
        run.font.bold = item_bold
    return box


def add_title(slide, eyebrow, title, subtitle=None, number=None):
    add_text(slide, eyebrow.upper(), 0.72, 0.45, 5.5, 0.25, 9, BLUE, True)
    add_text(slide, title, 0.72, 0.82, 11.3, 0.55, 27, INK, True)
    if subtitle:
        add_text(slide, subtitle, 0.72, 1.48, 11.4, 0.35, 12, MUTED)
    if number is not None:
        add_text(slide, f"{number:02d}", 12.15, 0.46, 0.5, 0.3, 10, MUTED, True, PP_ALIGN.RIGHT)


def add_footer(slide, text="SupportDesk · DevSecOps capstone"):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.72), Inches(7.08), Inches(11.85), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()
    add_text(slide, text, 0.72, 7.17, 6, 0.2, 8, MUTED)


def add_card(slide, x, y, w, h, title, body, accent=BLUE, fill=WHITE):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    card.line.color.rgb = LINE
    card.line.width = Pt(0.8)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.07), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_text(slide, title, x + 0.25, y + 0.2, w - 0.45, 0.28, 13, INK, True)
    add_text(slide, body, x + 0.25, y + 0.62, w - 0.45, h - 0.75, 10.5, SLATE)
    return card


def add_pill(slide, text, x, y, w, fill=BLUE_LIGHT, color=BLUE):
    pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.32))
    pill.fill.solid()
    pill.fill.fore_color.rgb = fill
    pill.line.fill.background()
    add_text(slide, text, x, y + 0.045, w, 0.18, 9, color, True, PP_ALIGN.CENTER)


def add_arrow(slide, x1, y1, x2, y2, color=BLUE):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = color
    line.line.width = Pt(2)
    line.line.end_arrowhead = True
    return line


def slide_one(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    add_text(slide, "DEVOPS FINAL CAPSTONE", 0.82, 0.75, 5, 0.25, 10, RGBColor(151, 184, 238), True)
    add_text(slide, "SupportDesk", 0.78, 1.35, 8.4, 0.8, 42, WHITE, True)
    add_text(slide, "A practical helpdesk ticket platform delivered through a complete DevSecOps pipeline.", 0.82, 2.35, 8.5, 0.7, 19, RGBColor(206, 216, 232))
    add_pill(slide, "React", 0.84, 3.55, 0.92, RGBColor(36, 55, 83), RGBColor(175, 207, 255))
    add_pill(slide, "FastAPI", 1.88, 3.55, 1.12, RGBColor(36, 55, 83), RGBColor(175, 207, 255))
    add_pill(slide, "PostgreSQL", 3.12, 3.55, 1.35, RGBColor(36, 55, 83), RGBColor(175, 207, 255))
    add_pill(slide, "Kubernetes", 4.58, 3.55, 1.2, RGBColor(36, 55, 83), RGBColor(175, 207, 255))
    add_pill(slide, "Terraform", 5.9, 3.55, 1.05, RGBColor(36, 55, 83), RGBColor(175, 207, 255))
    # Lifecycle visual
    add_text(slide, "Ticket lifecycle", 8.7, 1.35, 2.8, 0.3, 12, RGBColor(151, 184, 238), True)
    statuses = [("OPEN", BLUE), ("IN PROGRESS", AMBER), ("WAITING", PURPLE), ("RESOLVED", GREEN)]
    for index, (label, color) in enumerate(statuses):
        x = 8.72 + (index % 2) * 1.85
        y = 2.0 + (index // 2) * 1.25
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.62), Inches(0.62))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        add_text(slide, str(index + 1), x, y + 0.12, 0.62, 0.25, 15, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, label, x + 0.78, y + 0.17, 1.0, 0.22, 10, WHITE, True)
    add_text(slide, "Local validation: 12 backend tests passed · Docker Compose smoke test passed", 0.82, 6.62, 9.8, 0.25, 10, RGBColor(163, 181, 209))
    add_text(slide, "01", 12.15, 0.5, 0.5, 0.2, 10, RGBColor(132, 153, 184), True, PP_ALIGN.RIGHT)


def slide_two(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "01 · Product", "Support teams need one operational view", "SupportDesk turns incoming customer problems into visible, measurable work.", 2)
    add_card(slide, 0.75, 2.1, 3.65, 2.0, "Capture the issue", "Create a ticket with a subject, description, requester, category, priority, and owner.", BLUE, BLUE_LIGHT)
    add_card(slide, 4.72, 2.1, 3.65, 2.0, "Move the work", "Advance tickets through OPEN, IN_PROGRESS, WAITING, and RESOLVED states from the queue.", AMBER, AMBER_LIGHT)
    add_card(slide, 8.69, 2.1, 3.65, 2.0, "Observe the service", "Expose health, readiness, metrics, and queue statistics for operators and dashboards.", GREEN, GREEN_LIGHT)
    add_text(slide, "What makes this a capstone application", 0.75, 4.72, 5.5, 0.3, 16, INK, True)
    add_rich_text(slide, [
        ("01  ", "A real CRUD workflow instead of a blank starter page.", SLATE, False),
        ("02  ", "A responsive interface with loading and backend-error states.", SLATE, False),
        ("03  ", "A database-backed API that can travel from laptop to Kubernetes.", SLATE, False),
    ], 0.8, 5.2, 8.8, 1.2, 13)
    add_pill(slide, "Responsive UI", 9.55, 5.05, 1.35, BLUE_LIGHT, BLUE)
    add_pill(slide, "REST API", 11.0, 5.05, 1.05, GREEN_LIGHT, GREEN)
    add_pill(slide, "Metrics", 9.55, 5.5, 1.35, AMBER_LIGHT, AMBER)
    add_pill(slide, "Migrations", 11.0, 5.5, 1.05, PURPLE_LIGHT, PURPLE)
    add_footer(slide)


def slide_three(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "02 · Stack", "A deliberately small, production-shaped stack", "Each tool has a clear responsibility in the delivery path.", 3)
    cards = [
        (0.75, 2.05, "Frontend", "React · Vite · CSS\nNginx unprivileged runtime", BLUE, BLUE_LIGHT),
        (3.85, 2.05, "Backend", "Python · FastAPI · SQLAlchemy\nHealth, readiness, metrics", GREEN, GREEN_LIGHT),
        (6.95, 2.05, "Data", "PostgreSQL 16\nAlembic migrations", AMBER, AMBER_LIGHT),
        (10.05, 2.05, "Quality", "Pytest\nIsolated SQLite test database", PURPLE, PURPLE_LIGHT),
        (0.75, 4.55, "Packaging", "Dockerfiles\nDocker Compose", BLUE, BLUE_LIGHT),
        (3.85, 4.55, "Delivery", "GitHub Actions\nGHCR · Trivy", GREEN, GREEN_LIGHT),
        (6.95, 4.55, "Platform", "Terraform · AWS VPC/EKS\nKubernetes · Helm", AMBER, AMBER_LIGHT),
        (10.05, 4.55, "Operations", "Ingress · HPA\nPrometheus · Grafana · ArgoCD", PURPLE, PURPLE_LIGHT),
    ]
    for x, y, title, body, accent, fill in cards:
        add_card(slide, x, y, 2.8, 1.7, title, body, accent, fill)
    add_footer(slide)


def slide_four(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "03 · Architecture", "The browser reaches one application surface", "Nginx handles the frontend and `/api` proxy while FastAPI owns business logic and persistence.", 4)
    add_card(slide, 0.8, 2.45, 2.2, 1.25, "Browser", "Support agent\nhttp://localhost:3000", BLUE, BLUE_LIGHT)
    add_card(slide, 3.55, 2.45, 2.2, 1.25, "Nginx", "Static React app\n/api reverse proxy", PURPLE, PURPLE_LIGHT)
    add_card(slide, 6.3, 2.45, 2.2, 1.25, "FastAPI", "CRUD API\nhealth/readiness/metrics", GREEN, GREEN_LIGHT)
    add_card(slide, 9.05, 2.45, 2.2, 1.25, "PostgreSQL", "Tickets table\nAlembic migrations", AMBER, AMBER_LIGHT)
    add_arrow(slide, 3.0, 3.08, 3.5, 3.08)
    add_arrow(slide, 5.75, 3.08, 6.25, 3.08)
    add_arrow(slide, 8.5, 3.08, 9.0, 3.08)
    add_card(slide, 2.05, 4.65, 2.65, 1.15, "Delivery plane", "GitHub → Actions → Docker → GHCR", BLUE, PANEL)
    add_card(slide, 5.1, 4.65, 2.65, 1.15, "Infrastructure plane", "Terraform → VPC → EKS", AMBER, PANEL)
    add_card(slide, 8.15, 4.65, 2.65, 1.15, "Operations plane", "Helm → Ingress/HPA → monitoring", GREEN, PANEL)
    add_arrow(slide, 3.37, 4.55, 3.37, 3.82, MUTED)
    add_arrow(slide, 6.42, 4.55, 6.42, 3.82, MUTED)
    add_arrow(slide, 9.47, 4.55, 9.47, 3.82, MUTED)
    add_footer(slide)


def slide_five(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "04 · CI/CD", "Every main-branch image is traceable", "The commit SHA is carried from source code to the container registry and deployment.", 5)
    stages = [
        ("01", "Test", "pytest -v\nnpm ci && npm run build", BLUE, BLUE_LIGHT),
        ("02", "Build", "backend image\nfrontend image", PURPLE, PURPLE_LIGHT),
        ("03", "Scan", "Trivy\nHIGH/CRITICAL gate", AMBER, AMBER_LIGHT),
        ("04", "Publish", "GHCR\n<commit-sha> tags", GREEN, GREEN_LIGHT),
        ("05", "Deploy", "Helm\nEKS opt-in", BLUE, BLUE_LIGHT),
    ]
    for index, (number, title, body, accent, fill) in enumerate(stages):
        x = 0.75 + index * 2.48
        add_card(slide, x, 2.4, 2.0, 1.65, title, body, accent, fill)
        add_text(slide, number, x + 0.16, 2.08, 0.45, 0.2, 10, accent, True)
        if index < len(stages) - 1:
            add_arrow(slide, x + 2.06, 3.2, x + 2.4, 3.2, MUTED)
    add_text(slide, "Traceability example", 0.78, 4.85, 3.0, 0.25, 15, INK, True)
    add_text(slide, "commit A  →  backend:<sha-A>  +  frontend:<sha-A>  →  Helm release", 0.78, 5.28, 10.8, 0.35, 16, SLATE, False)
    add_text(slide, "A failed test or security gate stops promotion before an image reaches GHCR.", 0.78, 5.92, 10.8, 0.3, 12, MUTED)
    add_footer(slide)


def slide_six(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "05 · Platform", "Terraform creates the landing zone; Helm runs the workload", "The repository separates AWS infrastructure from Kubernetes application packaging.", 6)
    add_card(slide, 0.8, 2.1, 3.45, 2.2, "Terraform / AWS", "VPC\nTwo public + two private subnets\nNAT gateway\nEKS control plane\nManaged node group", AMBER, AMBER_LIGHT)
    add_card(slide, 4.65, 2.1, 3.45, 2.2, "Kubernetes", "supportdesk namespace\nFrontend + backend Deployments\nPostgreSQL + PVC\nClusterIP Services\nIngress + HPA", BLUE, BLUE_LIGHT)
    add_card(slide, 8.5, 2.1, 3.45, 2.2, "Helm release", "Reusable values\nSHA-based images\nHealth/readiness probes\nEnvironment-specific configuration", GREEN, GREEN_LIGHT)
    add_arrow(slide, 4.28, 3.2, 4.58, 3.2, MUTED)
    add_arrow(slide, 8.13, 3.2, 8.43, 3.2, MUTED)
    add_text(slide, "Operational prerequisites", 0.8, 4.92, 3.5, 0.28, 15, INK, True)
    add_rich_text(slide, [
        ("Ingress controller  ", "implements the nginx Ingress class.", SLATE, False),
        ("Metrics Server  ", "provides CPU data for the HPA.", SLATE, False),
        ("StorageClass  ", "provides persistent storage for PostgreSQL.", SLATE, False),
    ], 0.85, 5.35, 10.8, 1.1, 12)
    add_footer(slide)


def slide_seven(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide)
    add_title(slide, "06 · Observability", "Health and metrics are part of the application contract", "Operators can distinguish process health, database readiness, traffic, latency, and errors.", 7)
    add_card(slide, 0.8, 2.05, 2.75, 1.55, "/health", "Liveness\nReturns UP when the process is alive", GREEN, GREEN_LIGHT)
    add_card(slide, 3.8, 2.05, 2.75, 1.55, "/ready", "Readiness\nChecks database access before traffic", BLUE, BLUE_LIGHT)
    add_card(slide, 6.8, 2.05, 2.75, 1.55, "/metrics", "Prometheus\nRequest counters and latency histograms", PURPLE, PURPLE_LIGHT)
    add_card(slide, 9.8, 2.05, 2.75, 1.55, "Grafana", "Live panels\nRate, errors, p95 latency", AMBER, AMBER_LIGHT)
    add_arrow(slide, 3.58, 2.82, 3.75, 2.82, MUTED)
    add_arrow(slide, 6.58, 2.82, 6.75, 2.82, MUTED)
    add_arrow(slide, 9.58, 2.82, 9.75, 2.82, MUTED)
    add_text(slide, "Prometheus flow", 0.8, 4.55, 3, 0.3, 16, INK, True)
    steps = [("FastAPI", BLUE), ("ServiceMonitor", PURPLE), ("Prometheus", AMBER), ("Grafana", GREEN)]
    for index, (label, color) in enumerate(steps):
        x = 1.0 + index * 2.85
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(5.1), Inches(2.0), Inches(0.72))
        shape.fill.solid()
        shape.fill.fore_color.rgb = color
        shape.line.fill.background()
        add_text(slide, label, x, 5.32, 2.0, 0.2, 13, WHITE, True, PP_ALIGN.CENTER)
        if index < len(steps) - 1:
            add_arrow(slide, x + 2.05, 5.46, x + 2.72, 5.46, MUTED)
    add_footer(slide)


def slide_eight(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, NAVY)
    add_text(slide, "07 · SECURITY + RELIABILITY", 0.82, 0.62, 5.5, 0.25, 10, RGBColor(151, 184, 238), True)
    add_text(slide, "What the project demonstrates", 0.78, 1.15, 9.5, 0.55, 29, WHITE, True)
    add_text(slide, "A small application can still have disciplined delivery controls.", 0.82, 1.9, 9.4, 0.3, 14, RGBColor(206, 216, 232))
    items = [
        ("Non-root containers", "Backend UID 10001 and unprivileged Nginx runtime.", BLUE),
        ("Automated quality gates", "12 isolated API tests plus frontend production build.", GREEN),
        ("Supply-chain visibility", "Trivy scans both images before SHA-tagged publication.", AMBER),
        ("Operational feedback", "Probes, metrics, dashboards, HPA, and troubleshooting labs.", PURPLE),
    ]
    for index, (title, body, color) in enumerate(items):
        x = 0.82 + (index % 2) * 5.95
        y = 2.75 + (index // 2) * 1.45
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.38), Inches(0.38))
        circle.fill.solid()
        circle.fill.fore_color.rgb = color
        circle.line.fill.background()
        add_text(slide, "✓", x, y + 0.05, 0.38, 0.2, 12, WHITE, True, PP_ALIGN.CENTER)
        add_text(slide, title, x + 0.58, y, 4.8, 0.25, 14, WHITE, True)
        add_text(slide, body, x + 0.58, y + 0.4, 4.9, 0.35, 10.5, RGBColor(179, 194, 218))
    add_text(slide, "Next live proof", 0.82, 5.95, 2.2, 0.25, 13, RGBColor(151, 184, 238), True)
    add_text(slide, "AWS plan/apply → EKS → Helm → Prometheus/Grafana → ArgoCD → final walkthrough", 0.82, 6.32, 10.8, 0.28, 12, WHITE, False)
    add_text(slide, "08", 12.15, 0.5, 0.5, 0.2, 10, RGBColor(132, 153, 184), True, PP_ALIGN.RIGHT)


def build_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide_one(prs)
    slide_two(prs)
    slide_three(prs)
    slide_four(prs)
    slide_five(prs)
    slide_six(prs)
    slide_seven(prs)
    slide_eight(prs)
    prs.save(OUTPUT)
    print(f"Created {OUTPUT}")


if __name__ == "__main__":
    build_deck()
